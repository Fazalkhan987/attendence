import io

import pytest

from app.errors import AppError
from app.db import AppDatabase
from app.services import exporting
from app.services.attendance import AttendanceFilter, AttendanceService
from app.services.auth import AuthService
from datetime import datetime, timezone
from helpers import Clock, make_settings, make_services, person


def _auth(tmp_path, t, **kw):
    s = make_settings(tmp_path, **kw)
    return AuthService(AppDatabase(":memory:"), s, clock=lambda: t[0])


def test_login_logout_and_cookie_token_not_stored_plain(tmp_path):
    t = [1000.0]
    a = _auth(tmp_path, t)
    a.create_admin("alice", "a-long-passphrase", "Alice")
    token, user = a.login("alice", "a-long-passphrase", "ua")
    assert user["username"] == "alice" and a.user_for_token(token)["id"] == user["id"]
    assert token not in str(a.db.query("SELECT * FROM sessions")[0]["token_hash"])
    a.logout(token)
    assert a.user_for_token(token) is None and a.user_for_token(None) is None and a.user_for_token("bogus") is None


def test_wrong_password_unknown_user_same_message_and_lockout(tmp_path):
    t = [1000.0]
    a = _auth(tmp_path, t, max_failed_logins=3, lockout_minutes=10)
    a.create_admin("alice", "a-long-passphrase")
    msgs = set()
    for u in ("alice", "nobody"):
        with pytest.raises(AppError) as e:
            a.login(u, "wrong")
        msgs.add(e.value.message)
    assert len(msgs) == 1
    for _ in range(2):
        with pytest.raises(AppError):
            a.login("alice", "wrong")
    with pytest.raises(AppError) as e:
        a.login("alice", "a-long-passphrase")                      # locked: even the right password is refused
    assert e.value.status == 429
    t[0] += 11 * 60
    a.login("alice", "a-long-passphrase")


def test_session_expiry_and_sliding(tmp_path):
    t = [1000.0]
    a = _auth(tmp_path, t, session_ttl_hours=1)
    a.create_admin("alice", "a-long-passphrase")
    token, _ = a.login("alice", "a-long-passphrase")
    t[0] += 3000
    assert a.user_for_token(token)                                  # still valid, and slid forward
    t[0] += 3000
    assert a.user_for_token(token)                                  # would have expired without sliding
    t[0] += 3700
    assert a.user_for_token(token) is None


def test_admin_validation_and_password_change(tmp_path):
    t = [1000.0]
    a = _auth(tmp_path, t)
    for bad in ("ab", "bad name", "x" * 70):
        with pytest.raises(AppError):
            a.create_admin(bad, "a-long-passphrase")
    with pytest.raises(AppError):
        a.create_admin("alice", "short")
    u = a.create_admin("alice", "a-long-passphrase")
    with pytest.raises(AppError) as e:
        a.create_admin("ALICE", "a-long-passphrase")
    assert e.value.status == 409
    t1, _ = a.login("alice", "a-long-passphrase")
    t2, _ = a.login("alice", "a-long-passphrase")
    with pytest.raises(AppError):
        a.change_password(u["id"], "wrong", "another-long-pass", t1)
    a.change_password(u["id"], "a-long-passphrase", "another-long-pass", t1)
    assert a.user_for_token(t1) and a.user_for_token(t2) is None     # other sessions revoked
    a.login("alice", "another-long-pass")


# ------------------------------------------------------------------ attendance
def _att(tmp_path, clock=None, **kw):
    clock = clock or Clock()
    s = make_settings(tmp_path, **kw)
    names = {"S1": ("Ada", "CS"), "S2": ("Bob_%", "Math")}
    def lookup(pid):
        if pid not in names:
            from face_engine.exceptions import PersonNotFoundError
            raise PersonNotFoundError(pid)
        return names[pid]
    return AttendanceService(AppDatabase(":memory:"), s, lookup, clock), clock


def test_auto_mark_once_per_day_and_next_day_allowed(tmp_path):
    a, clock = _att(tmp_path)
    r1 = a.mark_from_face("S1", "Ada", "CS", 0.61)
    r2 = a.mark_from_face("S1", "Ada", "CS", 0.70)
    assert (r1["outcome"], r2["outcome"]) == ("created", "already_marked")
    assert a.list(AttendanceFilter())["total"] == 1 and r1["record"]["att_time"] == "08:30:00"
    clock.dt = clock.dt.replace(day=3)
    assert a.mark_from_face("S1", "Ada", "CS", 0.6)["outcome"] == "created"
    assert a.list(AttendanceFilter())["total"] == 2


def test_late_status_and_absent_upgrade(tmp_path):
    a, clock = _att(tmp_path, late_after="09:00")
    clock.dt = clock.dt.replace(hour=9, minute=30)
    assert a.mark_from_face("S1", "Ada", "CS", 0.6)["record"]["status"] == "late"
    a.create_manual("S2", "absent", None, None, None, "web:x")
    res = a.mark_from_face("S2", "Bob", "Math", 0.7)
    assert res["outcome"] == "updated" and res["record"]["status"] == "late" and res["record"]["source"] == "face"
    a.update(res["record"]["id"], "excused", None, "doctor", "web:x")
    assert a.mark_from_face("S2", "Bob", "Math", 0.7)["outcome"] == "already_marked"      # excused is not overridden


def test_manual_crud_and_validation(tmp_path):
    a, _ = _att(tmp_path)
    rec = a.create_manual("S1", "present", "2026-03-01", "07:45", "  note ", "web:x")
    assert rec["att_time"] == "07:45:00" and rec["source"] == "manual"
    with pytest.raises(AppError) as e:
        a.create_manual("S1", "present", "2026-03-01", None, None, "web:x")
    assert e.value.status == 409
    for args in (("S1", "weird", None, None), ("S1", "present", "2026-13-40", None), ("S1", "present", None, "25:99")):
        with pytest.raises(AppError):
            a.create_manual(*args, None, "web:x")
    from face_engine.exceptions import PersonNotFoundError
    with pytest.raises(PersonNotFoundError):
        a.create_manual("NOPE", "present", None, None, None, "web:x")
    assert a.update(rec["id"], "late", "08:00", None, "web:y")["status"] == "late"
    a.delete(rec["id"])
    with pytest.raises(AppError):
        a.delete(rec["id"])


def test_filters_search_escaping_pagination(tmp_path):
    a, _ = _att(tmp_path)
    for d in ("2026-03-01", "2026-03-02", "2026-03-03"):
        a.create_manual("S1", "present", d, None, None, "x")
    a.create_manual("S2", "absent", "2026-03-02", None, None, "x")
    assert a.list(AttendanceFilter(date_from="2026-03-02", date_to="2026-03-02"))["total"] == 2
    assert a.list(AttendanceFilter(status="absent"))["total"] == 1
    assert a.list(AttendanceFilter(q="math"))["total"] == 1                       # department search
    assert a.list(AttendanceFilter(q="%"))["total"] == 1                          # '%' is literal, matches only "Bob_%"
    assert a.list(AttendanceFilter(q="_"))["total"] == 1
    p = a.list(AttendanceFilter(), page=2, page_size=3)
    assert p["total"] == 4 and len(p["items"]) == 1
    assert a.list(AttendanceFilter())["items"][0]["att_date"] == "2026-03-03"      # newest first
    with pytest.raises(AppError):
        a.list(AttendanceFilter(date_from="2026-03-05", date_to="2026-03-01"))


def test_counts_trend_sync_delete(tmp_path):
    a, _ = _att(tmp_path)
    a.create_manual("S1", "present", "2026-03-02", None, None, "x")
    a.create_manual("S2", "late", "2026-03-02", None, None, "x")
    assert a.counts_for_date("2026-03-02")["late"] == 1 and a.present_ids_for_date("2026-03-02") == {"S1", "S2"}
    tr = a.trend(3)
    assert [d["date"] for d in tr] == ["2026-02-28", "2026-03-01", "2026-03-02"] and tr[-1]["present"] == 1 and tr[-1]["late"] == 1
    a.sync_person("S1", "Ada L.", "Math")
    assert a.list(AttendanceFilter(person_id="S1"))["items"][0]["person_name"] == "Ada L."
    assert a.delete_for_person("S1") == 1


def test_export_csv_injection_and_xlsx(tmp_path):
    rows = [{"att_date": "2026-03-02", "att_time": "08:00:00", "person_id": "S1", "person_name": "=HYPERLINK(\"http://x\")",
             "department": "+cmd", "status": "present", "source": "face", "confidence": 0.61234, "note": None, "marked_by": "@a"}]
    csv_text = exporting.to_csv(rows).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text and "'+cmd" in csv_text and "'@a" in csv_text and "0.612" in csv_text
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(exporting.to_xlsx(rows))).active
    assert ws["A1"].value == "Date" and ws["D2"].value.startswith("'=") and ws.max_row == 2
